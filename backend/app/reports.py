"""Report builders: CSV, Excel (.xlsx) and PDF from classified rows.

Pure functions that build everything in memory (io.StringIO / io.BytesIO). They never
write files or temp files, and never log row data.

Spreadsheet formula injection: any text cell that starts with = + - @, a tab or a
carriage return gets a leading ' so Excel and LibreOffice treat it as text. Numbers
stay numeric, so negative amounts are not touched.

Money in (rows with direction CR, from bank PDFs) is not spending: the category totals
leave it out, and each report shows it on its own "Money in" line when there is any.

The Excel and PDF reports also carry the insights from classifier.insights (an "Insights"
sheet and a "Highlights" section). The CSV report does not.
"""

import csv
import io
from datetime import date
from xml.sax.saxutils import escape

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import PieChart, Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.legends import Legend
from reportlab.graphics.charts.piecharts import Pie
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.lib.enums import TA_RIGHT
from reportlab.platypus import CondPageBreak, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.schemas import Row, Summary
from classifier.insights import compute_insights

FIELDS = list(Row.model_fields)
TITLE = "Spending report"
FOOTER = "Generated locally. No data left this computer."
# 50/30/20 PDF table columns in mm; the A4 content width is 180 mm (15 mm margins).
BUDGET_COL_MM = [62, 40, 30, 30]

# Must match CATEGORY_COLORS and OTHER_COLOR in frontend/src/lib/format.ts (a test checks this).
CATEGORY_COLORS = {
    "food": "#D85A30",
    "transport": "#378ADD",
    "shopping": "#D4537E",
    "bills": "#7F77DD",
    "transfer": "#1D9E75",
}
OTHER_COLOR = "#888780"

_INJECTION_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def safe_text(value: str) -> str:
    """Prefix ' to text a spreadsheet would read as a formula."""
    return "'" + value if value.startswith(_INJECTION_PREFIXES) else value


def _safe(value):
    return safe_text(value) if isinstance(value, str) else value


def format_idr(value: float) -> str:
    """Rp 6.731.990 style, no decimals."""
    sign = "-" if value < 0 else ""
    return f"{sign}Rp {round(abs(value)):,}".replace(",", ".")


def period(rows: list[Row]) -> str | None:
    """'YYYY-MM-DD to YYYY-MM-DD' from the rows' dates, or None if none parse."""
    dates = [r.date for r in rows if r.date]
    if not dates:
        return None
    raw = pd.Series(dates)
    # ISO dates first; anything else (e.g. 18/09/2026 from a bank export) is read day-first.
    parsed = pd.to_datetime(raw, errors="coerce", format="ISO8601")
    rest = parsed.isna()
    if rest.any():
        parsed[rest] = pd.to_datetime(raw[rest], errors="coerce", format="mixed", dayfirst=True)
    parsed = parsed.dropna()
    if parsed.empty:
        return None
    return f"{parsed.min().date().isoformat()} to {parsed.max().date().isoformat()}"


def _color(category: str) -> str:
    return CATEGORY_COLORS.get(category, OTHER_COLOR)


# ---- CSV -------------------------------------------------------------------


def report_csv(rows: list[Row], summary: Summary) -> str:
    """Summary block, a blank line, then the row table."""
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(["category", "count", "total_amount", "share"])
    for c in summary.by_category:
        writer.writerow([_safe(c.category), c.count, c.total_amount, c.share])
    if summary.money_in_count:
        writer.writerow(["money in", summary.money_in_count, summary.money_in_total, ""])
    writer.writerow([])
    writer.writerow(FIELDS)
    for r in rows:
        writer.writerow(["" if v is None else _safe(v) for v in r.model_dump().values()])
    return buf.getvalue()


# ---- Excel -----------------------------------------------------------------


def _xl_text(value: str) -> str:
    # openpyxl refuses control characters that XML cannot hold.
    return safe_text(ILLEGAL_CHARACTERS_RE.sub("", value))


def _xl(value):
    return _xl_text(value) if isinstance(value, str) else value


_TX_WIDTHS = {"id": 6, "date": 12, "description": 42, "amount": 14, "direction": 10, "predicted": 12,
              "confidence": 12, "needs_review": 14, "source": 9}


def _insights_sheet(wb: Workbook, insights: dict) -> None:
    ws = wb.create_sheet("Insights")
    bold = Font(bold=True)
    idr = "#,##0"
    pct = "0.0%"
    ws.append(["Insights"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([_xl_text(insights["disclaimer"])])

    def section(title: str, header: list[str] | None = None) -> None:
        ws.append([])
        ws.append([title])
        ws.cell(row=ws.max_row, column=1).font = bold
        if header:
            ws.append(header)
            for cell in ws[ws.max_row]:
                cell.font = bold

    def add(values: list, formats: list[str | None]) -> None:
        ws.append([_xl(v) for v in values])
        for col, fmt in enumerate(formats, start=1):
            if fmt:
                ws.cell(row=ws.max_row, column=col).number_format = fmt

    cash = insights["cash_flow"]
    section("Cash flow")
    for label, key in [("Money in (IDR)", "money_in"), ("Total out (IDR)", "total_out"),
                       ("Living spend (IDR)", "living_spend"), ("Transfers out (IDR)", "transfers_out"),
                       ("Net (IDR)", "net")]:
        add([label, cash[key]], [None, idr])
    add(["Savings rate", cash["savings_rate"] if cash["savings_rate"] is not None else "-"], [None, pct])

    budget = insights["budget_503020"]
    basis = "money in" if budget["basis"] == "money_in" else "total out (no money in found)"
    section("50/30/20 view", ["Bucket", "Amount (IDR)", "Share", "Guideline"])
    for label, key in [("Needs (bills + transport)", "needs"), ("Wants (food + shopping)", "wants"),
                       ("Savings (net)", "savings"), ("Transfers (unclassified)", "unclassified_transfers")]:
        amount = budget[key]
        share = budget["shares"][key]
        target = budget["targets"].get(key)
        add([label, amount if amount is not None else "-", share if share is not None else "-",
             target if target is not None else "-"], [None, idr, pct, pct])
    add([f"Shares of {basis}"], [None])

    section("Top merchants by total", ["Merchant", "Count", "Total (IDR)", "Average (IDR)"])
    for m in insights["top_merchants"]["by_total"]:
        add([m["merchant"], m["count"], m["total"], m["avg"]], [None, None, idr, idr])

    small = insights["small_spends"]
    section("Small spends")
    add([f"Under {format_idr(small['threshold'])}: count", small["count"]], [None, None])
    add(["Small spends total (IDR)", small["total"]], [None, idr])
    add(["Share of living spend", small["share_of_living"] if small["share_of_living"] is not None else "-"],
        [None, pct])

    fees = insights["fees"]
    section("Fees")
    add(["Fee count", fees["count"]], [None, None])
    add(["Fees total (IDR)", fees["total"]], [None, idr])

    section("Highlights", ["Highlight", "Impact (IDR)"])
    for h in insights["highlights"]:
        add([h["text"], h["impact_rp"]], [None, idr])

    for col, width in zip("ABCD", (60, 16, 16, 16)):
        ws.column_dimensions[col].width = width


def report_xlsx(rows: list[Row], summary: Summary, insights: dict | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    bold = Font(bold=True)

    ws["A1"] = TITLE
    ws["A1"].font = Font(bold=True, size=14)
    info = [
        ("Generated", date.today().isoformat()),
        ("Period", period(rows) or "-"),
        ("Total spend (IDR)", summary.total_amount),
        ("Transactions", summary.total_rows),
        ("Rows needing review", summary.needs_review),
    ]
    if summary.money_in_count:
        info += [("Money in (IDR)", summary.money_in_total), ("Money-in transactions", summary.money_in_count)]
    for i, (label, value) in enumerate(info, start=3):
        ws.cell(row=i, column=1, value=label).font = bold
        cell = ws.cell(row=i, column=2, value=_xl(value))
        if label.endswith("(IDR)"):
            cell.number_format = "#,##0"

    header_row = 3 + len(info) + 1
    for col, name in enumerate(["Category", "Count", "Total (IDR)", "Share"], start=1):
        ws.cell(row=header_row, column=col, value=name).font = bold
    for i, c in enumerate(summary.by_category, start=header_row + 1):
        ws.cell(row=i, column=1, value=_xl(c.category))
        ws.cell(row=i, column=2, value=c.count)
        ws.cell(row=i, column=3, value=c.total_amount).number_format = "#,##0"
        ws.cell(row=i, column=4, value=c.share).number_format = "0.0%"
    last_row = header_row + len(summary.by_category)
    for col, width in zip("ABCD", (22, 12, 16, 10)):
        ws.column_dimensions[col].width = width

    if summary.by_category:
        # Pie of totals by category; by count when the upload has no amounts.
        value_col = 3 if summary.total_amount > 0 else 2
        chart = PieChart()
        chart.title = "Spend by category" if value_col == 3 else "Transactions by category"
        chart.add_data(Reference(ws, min_col=value_col, min_row=header_row, max_row=last_row), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=header_row + 1, max_row=last_row))
        chart.height, chart.width = 9, 12
        ws.add_chart(chart, "F3")

    tx = wb.create_sheet("Transactions")
    tx.append(FIELDS)
    for cell in tx[1]:
        cell.font = bold
    for r in rows:
        tx.append([_xl(v) for v in r.model_dump().values()])
    for col, field in enumerate(FIELDS, start=1):
        letter = get_column_letter(col)
        if field == "amount":
            for (cell,) in tx.iter_rows(min_row=2, min_col=col, max_col=col):
                cell.number_format = "#,##0"
        elif field == "confidence":
            for (cell,) in tx.iter_rows(min_row=2, min_col=col, max_col=col):
                cell.number_format = "0.00"
        tx.column_dimensions[letter].width = _TX_WIDTHS.get(field, 12)
    tx.freeze_panes = "A2"
    tx.auto_filter.ref = tx.dimensions

    _insights_sheet(wb, insights if insights is not None else compute_insights(rows))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- PDF -------------------------------------------------------------------


def _pie(summary: Summary) -> Drawing | None:
    by_amount = summary.total_amount > 0
    slices = [(c.category, c.total_amount if by_amount else c.count) for c in summary.by_category]
    slices = [(name, value) for name, value in slices if value > 0]
    if not slices:
        return None

    drawing = Drawing(170 * mm, 62 * mm)
    drawing.initialFontName = "Helvetica"  # reportlab's drawing default is Times-Roman
    pie = Pie()
    pie.x, pie.y = 5 * mm, 3 * mm
    pie.width = pie.height = 56 * mm
    pie.data = [value for _, value in slices]
    pie.simpleLabels = 1
    pie.labels = None
    pie.slices.strokeColor = colors.white
    pie.slices.strokeWidth = 1
    for i, (name, _) in enumerate(slices):
        pie.slices[i].fillColor = colors.HexColor(_color(name))
    drawing.add(pie)

    total = sum(value for _, value in slices)
    legend = Legend()
    legend.x, legend.y = 75 * mm, 50 * mm
    legend.alignment = "right"
    legend.fontName = "Helvetica"
    legend.fontSize = 9
    legend.columnMaximum = 10
    legend.dy = legend.dx = 8
    legend.deltay = 14
    legend.colorNamePairs = [
        (colors.HexColor(_color(name)), f"{name.capitalize()}  {value / total:.1%}") for name, value in slices
    ]
    drawing.add(legend)
    return drawing


def _footer(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#64748b"))
    canvas.drawString(doc.leftMargin, 10 * mm, FOOTER)
    canvas.drawRightString(A4[0] - doc.rightMargin, 10 * mm, f"Page {doc.page}")
    canvas.restoreState()


_GRID = [
    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
    ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
    ("FONTSIZE", (0, 0), (-1, -1), 8),
    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
    ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
    ("VALIGN", (0, 0), (-1, -1), "TOP"),
]


def _table(data: list[list], widths_mm: list[float], cell, right_from: int = 1) -> Table:
    """A grid table with explicit column widths and a header row that repeats on every page.

    Header labels become bold Paragraphs so a long label wraps inside its column instead of
    running into the next one. Columns from `right_from` on are right-aligned.
    """
    head = cell.clone("head", fontName="Helvetica-Bold")
    head_right = head.clone("head_right", alignment=TA_RIGHT)
    header = [Paragraph(escape(str(h)), head_right if i >= right_from else head) for i, h in enumerate(data[0])]
    t = Table([header] + data[1:], colWidths=[w * mm for w in widths_mm], repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle(_GRID + [("ALIGN", (right_from, 0), (-1, -1), "RIGHT")]))
    return t


def _budget_table(budget: dict, cell) -> list:
    """The 50/30/20 table (short headers, widths that fit A4) plus a caption naming the share basis."""
    basis = "money in" if budget["basis"] == "money_in" else "total out (no money in found)"

    def money(value):
        return format_idr(value) if value is not None else "-"

    def share(value):
        return f"{value:.1%}" if value is not None else "-"

    table = [["50/30/20 view", "Amount", "Share", "Guideline"]]
    for label, key in [("Needs (bills + transport)", "needs"), ("Wants (food + shopping)", "wants"),
                       ("Savings (net)", "savings"), ("Transfers (unclassified)", "unclassified_transfers")]:
        table.append([label, money(budget[key]), share(budget["shares"][key]), share(budget["targets"].get(key))])
    caption = cell.clone("caption", fontSize=7, leading=9, textColor=colors.HexColor("#64748b"))
    return [_table(table, BUDGET_COL_MM, cell), Spacer(1, 1 * mm), Paragraph(f"Share = share of {basis}.", caption)]


def _highlights_story(insights: dict, styles, cell) -> list:
    body = styles["BodyText"]
    if insights["highlights"]:
        items = [Paragraph(f"{i}. {escape(h['text'])}", body) for i, h in enumerate(insights["highlights"], start=1)]
    else:
        items = [Paragraph("Not enough amounts in this statement for highlights.", body)]
    story = [KeepTogether([Paragraph("Highlights", styles["Heading2"]), *items]), Spacer(1, 3 * mm)]

    # KeepTogether: a table header never sits alone at the bottom of a page.
    story += [KeepTogether(_budget_table(insights["budget_503020"], cell)), Spacer(1, 3 * mm)]

    merchants = insights["top_merchants"]["by_total"]
    if merchants:
        rows = [["Top merchants by total", "Count", "Total", "Average"]] + [
            [Paragraph(escape(m["merchant"]), cell), str(m["count"]), format_idr(m["total"]), format_idr(m["avg"])]
            for m in merchants
        ]
        story += [KeepTogether(_table(rows, [70, 20, 35, 32], cell)), Spacer(1, 2 * mm)]

    story += [Paragraph(escape(insights["disclaimer"]), cell), Spacer(1, 6 * mm)]
    return story


def report_pdf(rows: list[Row], summary: Summary, insights: dict | None = None) -> bytes:
    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    cell = body.clone("cell", fontName="Helvetica", fontSize=8, leading=10)
    has_amounts = summary.total_amount > 0

    story = [
        Paragraph(TITLE, styles["Title"]),
        Paragraph(f"Generated {date.today().isoformat()}", body),
        Paragraph(f"Period: {escape(period(rows) or '-')}", body),
        Paragraph(f"Total spend: <b>{format_idr(summary.total_amount) if has_amounts else '-'}</b>", body),
        Paragraph(f"Transactions: {summary.total_rows:,}".replace(",", "."), body),
        Paragraph(f"Needs review: {summary.needs_review:,}".replace(",", "."), body),
    ]
    if summary.money_in_count:
        story.append(Paragraph(
            f"Money in: {format_idr(summary.money_in_total)} "
            f"({summary.money_in_count} transactions, not counted as spend)", body))
    story.append(Spacer(1, 6 * mm))

    pie = _pie(summary)
    if pie is not None:
        heading = "Share of spend by category" if has_amounts else "Share of transactions by category"
        story += [KeepTogether([Paragraph(heading, styles["Heading2"]), pie]), Spacer(1, 4 * mm)]

    totals = [["Category", "Count", "Total", "Share"]] + [
        [c.category.capitalize(), str(c.count), format_idr(c.total_amount) if has_amounts else "-", f"{c.share:.1%}"]
        for c in summary.by_category
    ]
    totals_table = _table(totals, [45, 20, 40, 20], cell)
    story += [KeepTogether([Paragraph("Totals by category", styles["Heading2"]), totals_table]), Spacer(1, 6 * mm)]
    story += _highlights_story(insights if insights is not None else compute_insights(rows), styles, cell)

    tx = [["Date", "Description", "Amount", "Category", "Source"]] + [
        [
            r.date or "",
            Paragraph(escape(r.description), cell),
            ("+" if r.direction == "CR" else "") + format_idr(r.amount) if r.amount is not None else "",
            r.predicted + (" *" if r.needs_review else ""),
            r.source,
        ]
        for r in rows
    ]
    tx_table = Table(tx, colWidths=[22 * mm, 80 * mm, 28 * mm, 24 * mm, 16 * mm], repeatRows=1, hAlign="LEFT")
    tx_table.setStyle(TableStyle(_GRID + [("ALIGN", (2, 0), (2, -1), "RIGHT")]))
    story += [
        # Start a new page unless the heading, note, header row and a few rows fit here.
        CondPageBreak(45 * mm),
        Paragraph("Transactions", styles["Heading2"]),
        Paragraph("* needs review" + ("; + money in" if summary.money_in_count else ""), cell),
        Spacer(1, 2 * mm),
        tx_table,
    ]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, title=TITLE, author="", creator="", subject="",
        leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=18 * mm,
    )
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()
