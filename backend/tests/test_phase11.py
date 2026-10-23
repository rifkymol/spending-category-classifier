"""Phase 11: label migration, the recipient memory, the 50/30/20 v2 math, transfers per category,
and the group subtotals and "Other money out" row in the reports.

Synthetic rows only: every merchant and person below is made up.
"""

import io
import logging

import joblib
import pandas as pd
import pdfplumber
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import classify as cli
from app.main import app
from classifier import MEMORY_ATTR, build_model, build_recipient_memory, classify, clean_description, recipient_memory
from classifier.insights import compute_insights
from test_insights import tx
from test_pipeline import SAMPLE, _labels_csv


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


# ---- migration -------------------------------------------------------------


def old_labels(path):
    return _labels_csv(path, [
        ["2026-09-01", "QR WARUNG CONTOH A", 15000, "DB", "food", "user"],
        ["2026-09-02", "QR WARUNG CONTOH B", 16000, "DB", "food", "user"],
        ["2026-09-03", "TRSF E-BANKING DB 0309 BUDI CONTOH", 250000, "DB", "transfer", "rule"],
        ["2026-09-04", "QR TOKO CONTOH", 20000, "DB", "shopping", "user"],
        ["2026-09-05", "PLN PRABAYAR", 100000, "DB", "bills", "rule"],
        ["2026-09-06", "GRAB CAR CONTOH", 25000, "DB", "transport", "model"],
    ])


def test_migrate_labels_maps_keys_and_marks_changed_ones():
    df = pd.DataFrame({"description": ["a", "b", "c", "d"], "category": ["food", "transfer", "bills", "groceries"],
                       "label_source": ["user", "rule", "user", "user"]})
    out = cli.migrate_labels(df)
    assert list(out["category"]) == ["eating_out", "other_transfer", "bills", "groceries"]
    assert list(out["label_source"]) == ["migrated", "migrated", "user", "user"]
    # Files without a label_source column get one only for the migrated rows.
    out = cli.migrate_labels(df.drop(columns=["label_source"]))
    assert list(out["label_source"]) == ["migrated", "migrated", None, None]


def test_train_excludes_migrated_rows(tmp_path, capsys):
    path = old_labels(tmp_path / "labels-old.csv")
    assert cli.main(["train", str(path), "--no-synthetic", "--model", str(tmp_path / "m.joblib")]) == 0
    out = capsys.readouterr().out
    assert "Dropped 1 unreviewed rows" in out
    assert "Excluded 3 migrated rows" in out
    assert "Trained on 2 rows (2 categories)" in out
    assert "  eating_out" not in out and "WARUNG" not in out.upper()  # migrated rows are not trained on


def test_evaluate_loads_old_key_files(tmp_path, capsys):
    path = old_labels(tmp_path / "labels-old.csv")
    assert cli.main(["evaluate", str(path), str(SAMPLE)]) == 0
    out = capsys.readouterr().out
    assert "Excluded 3 migrated rows" in out
    assert "food" not in out.split("precision", 1)[1]  # the report lists new keys only


def test_old_model_predictions_are_mapped_and_flagged():
    """A model trained before Phase 11 still predicts food/transfer: map them and ask for review."""
    old = build_model().fit(["makan contoh", "barang contoh", "orang contoh"], ["food", "shopping", "transfer"])
    out = classify(pd.DataFrame({"description": ["MAKAN CONTOH", "ORANG CONTOH", "BARANG CONTOH"]}), old, threshold=0.0)
    assert list(out["model_category"]) == ["eating_out", "other_transfer", "shopping"]
    assert list(out["needs_review"]) == [True, True, False]


# ---- recipient memory ------------------------------------------------------


def memory_labels(path):
    return _labels_csv(path, [
        ["2026-09-01", "TRSF E-BANKING DB 0109/FTSCY/WS1 IBU CONTOH", 500000, "DB", "people_transfer", "user"],
        ["2026-09-10", "TRSF E-BANKING DB 1009/FTSCY/WS2 IBU CONTOH", 500000, "DB", "people_transfer", "user"],
        ["2026-09-11", "TRSF E-BANKING DB 1109/FTSCY/WS3 IBU CONTOH", 500000, "DB", "other_transfer", "rule"],
        ["2026-09-05", "TRSF E-BANKING DB 0509/FTSCY/WS4 REKENING SAYA CONTOH", 2000000, "DB", "savings_transfer",
         "user"],
        ["2026-09-06", "TRSF E-BANKING DB 0609/FTSCY/WS5 KFC SUDIRMAN", 50000, "DB", "people_transfer", "user"],
        ["2026-09-07", "TRSF E-BANKING DB 0709/FTSCY/WS6 TEMAN LAMA", 90000, "DB", "people_transfer", "model"],
        ["2026-09-08", "TARIK TUNAI ATM CONTOH", 300000, "DB", "cash", "rule"],
        ["2026-09-09", "QR WARUNG CONTOH", 20000, "DB", "eating_out", "user"],
    ])


def test_build_recipient_memory():
    df = pd.read_csv(io.StringIO(
        "description,category,label_source\n"
        "TRSF KE IBU CONTOH,people_transfer,user\n"
        "TRSF KE IBU CONTOH,people_transfer,rule\n"
        "TRSF KE IBU CONTOH,other_transfer,user\n"
        "TRSF KE ADIK CONTOH,people_transfer,model\n"
        "TRSF KE ADIK CONTOH,cash,memory\n"
        "QR WARUNG CONTOH,eating_out,user\n"
        "BI-FAST TRANSFER REF 91,other_transfer,user\n"
        "TARIK TUNAI ATM 01,cash,user\n"))
    # Most frequent wins; model/memory rows, non-transfer rows and unnamed counterparties are left out.
    assert build_recipient_memory(df) == {"Ibu Contoh": "people_transfer"}
    assert build_recipient_memory(df.drop(columns=["label_source"])) == {}


def test_train_saves_memory_and_classify_uses_it(tmp_path, capsys):
    model_path = tmp_path / "personal.joblib"
    assert cli.main(["train", str(memory_labels(tmp_path / "labels.csv")), "--model", str(model_path)]) == 0
    out = capsys.readouterr().out
    assert "Recipient memory: 3 recipients" in out
    assert "IBU" not in out.upper() and "REKENING" not in out.upper()  # counts only

    model = joblib.load(model_path)
    # The synthetic sample (no label_source) adds nothing; the "model" row is not a label.
    assert recipient_memory(model) == {"Ibu Contoh": "people_transfer", "Rekening Saya Contoh": "savings_transfer",
                                       "Kfc Sudirman": "people_transfer"}

    new = pd.DataFrame({"description": [
        "TRSF E-BANKING DB 2109/FTSCY/WS9 IBU CONTOH",
        "BI-FAST DB TRANSFER REKENING SAYA CONTOH",
        "TRSF E-BANKING DB 2209/FTSCY/WS8 KFC SUDIRMAN",  # a non-transfer rule beats the memory
        "TRSF E-BANKING DB 2309/FTSCY/WS7 TEMAN LAMA",  # not in the memory: the transfer rule
    ]})
    result = classify(new, model)
    assert list(result["predicted"]) == ["people_transfer", "savings_transfer", "eating_out", "other_transfer"]
    assert list(result["source"]) == ["memory", "memory", "rule", "rule"]
    assert not result["needs_review"].any()


def test_memory_applies_to_model_transfer_rows_only():
    model = build_model().fit(["orang satu", "orang dua", "toko satu", "toko dua"],
                              ["people_transfer", "people_transfer", "shopping", "shopping"])
    setattr(model, MEMORY_ATTR, {"Orang Satu": "savings_transfer", "Toko Satu": "savings_transfer"})
    out = classify(pd.DataFrame({"description": ["ORANG SATU", "TOKO SATU"]}), model)
    assert list(out["predicted"]) == ["savings_transfer", "shopping"]
    assert list(out["source"]) == ["memory", "model"]


def test_synthetic_model_has_empty_memory(client):
    assert recipient_memory(app.state.model) == {}
    rows = client.post("/api/classify", content=SAMPLE.read_bytes(), headers={"Content-Type": "text/csv"}).json()["rows"]
    assert not any(r["source"] == "memory" for r in rows)


def test_api_shows_memory_source(client, tmp_path, monkeypatch, caplog):
    model_path = tmp_path / "personal.joblib"
    assert cli.main(["train", str(memory_labels(tmp_path / "labels.csv")), "--model", str(model_path)]) == 0
    monkeypatch.setenv("MODEL_PATH", str(model_path))
    saved = app.state.model, app.state.model_kind
    caplog.set_level(logging.DEBUG)
    try:
        with TestClient(app) as custom:
            csv = b"description,amount\nTRSF E-BANKING DB 2109/FTSCY/WS9 IBU CONTOH,750000\n"
            rows = custom.post("/api/classify", content=csv, headers={"Content-Type": "text/csv"}).json()["rows"]
    finally:
        app.state.model, app.state.model_kind = saved
    assert rows[0]["predicted"] == "people_transfer" and rows[0]["source"] == "memory"
    assert rows[0]["confidence"] == 1.0 and rows[0]["needs_review"] is False
    text = "\n".join(r.getMessage() for r in caplog.records)
    assert "IBU" not in text.upper() and "750000" not in text


# ---- 50/30/20 v2 -----------------------------------------------------------


def month():
    """One month: income 10.000.000; needs 3.500.000; wants 1.000.000; savings_transfer 1.500.000;
    people 700.000 + cash 300.000 = other money out 1.000.000. Total out 7.000.000."""
    rows = [
        tx("KR OTOMATIS GAJI PT CONTOH", 10_000_000, "other_transfer", "CR", "2026-09-01"),
        tx("INDOMARET CONTOH", 1_000_000, "groceries", date="2026-09-02"),
        tx("PLN PRABAYAR", 500_000, "bills", date="2026-09-03"),
        tx("TRSF E-BANKING DB PAK KOS CONTOH", 2_000_000, "housing", date="2026-09-04"),
        tx("WARUNG CONTOH", 800_000, "eating_out", date="2026-09-05"),
        tx("NETFLIX CONTOH", 200_000, "entertainment", date="2026-09-06"),
        tx("TOPUP DANA CONTOH", 1_500_000, "savings_transfer", date="2026-09-07"),
        tx("TRSF KE IBU CONTOH", 700_000, "people_transfer", date="2026-09-08"),
        tx("TARIK TUNAI ATM 01", 300_000, "cash", date="2026-09-30"),
    ]
    for i, r in enumerate(rows):
        r["id"] = i
    return rows


def test_budget_v2_by_hand():
    ins = compute_insights(month())
    b = ins["budget_503020"]
    assert (b["needs"], b["wants"], b["savings_transfers"], b["other_out"]) == (3_500_000, 1_000_000, 1_500_000, 1_000_000)
    # Savings = income - spending, and spending leaves out the 1.500.000 moved to own savings:
    # 10.000.000 - (7.000.000 - 1.500.000) = 4.500.000 (the net 3.000.000 plus the savings transfer).
    assert b["savings"] == 4_500_000
    assert ins["cash_flow"]["left_over"] == 3_000_000 and ins["cash_flow"]["saved"] == 4_500_000
    assert b["actual_amounts"] == {"needs": 3_500_000, "wants": 1_000_000, "savings": 4_500_000}
    assert b["target_amounts"] == {"needs": 5_000_000, "wants": 3_000_000, "savings": 2_000_000}
    assert b["difference"] == {"needs": 1_500_000, "wants": 2_000_000, "savings": -2_500_000}
    assert b["shares"] == {"needs": 0.35, "wants": 0.1, "savings": 0.45, "savings_transfers": 0.15, "other_out": 0.1}
    assert b["savings_transfers_monthly"] == 1_500_000 and b["other_out_monthly"] == 1_000_000
    # Cash flow: the whole Transfers group is "transfers out"; living spend is Needs + Wants.
    assert ins["cash_flow"]["transfers_out"] == 2_500_000 and ins["cash_flow"]["living_spend"] == 4_500_000


def test_savings_transfer_counts_as_savings():
    """Moving money to your own savings raises savings by that amount; giving it to people does not."""
    base = month()
    as_people = [dict(r, predicted="people_transfer") if r["predicted"] == "savings_transfer" else r for r in base]
    saved = compute_insights(base)["budget_503020"]["actual_amounts"]["savings"]
    given = compute_insights(as_people)["budget_503020"]["actual_amounts"]["savings"]
    assert saved - given == 1_500_000


def test_budget_v2_without_income():
    rows = [dict(r, direction=None) for r in month() if r["direction"] != "CR"]
    b = compute_insights(rows)["budget_503020"]
    assert b["basis"] == "total_out" and b["basis_amount"] == 7_000_000
    assert b["savings"] is None and b["actual_amounts"]["savings"] is None and b["difference"] is None
    assert b["shares"]["savings_transfers"] == pytest.approx(1_500_000 / 7_000_000, abs=1e-4)
    assert b["shares"]["other_out"] == pytest.approx(1_000_000 / 7_000_000, abs=1e-4)


def test_budget_v2_with_override():
    b = compute_insights(month(), income_override=6_000_000)["budget_503020"]
    assert b["income_source"] == "override"
    assert b["actual_amounts"]["savings"] == 6_000_000 - 5_500_000
    assert b["difference"]["savings"] == 1_200_000 - 500_000


def test_v2_highlights_in_both_languages():
    other = [tx("TRSF KE IBU CONTOH", 700_000, "people_transfer"), tx("TARIK TUNAI ATM", 300_000, "cash")]
    texts = [h["text"] for h in compute_insights(other)["highlights"]]
    assert any(x.startswith("Rp 1.000.000 went to family and friends, cash and other transfers.") for x in texts)
    texts = [h["text"] for h in compute_insights(other, lang="id")["highlights"]]
    assert any(x.startswith("Rp 1.000.000 keluar ke keluarga dan teman") for x in texts)
    # Without money in (a CSV upload) the move to savings gets its own highlight.
    csv = [dict(r, direction=None) for r in month() if r["direction"] != "CR"]
    texts = [h["text"] for h in compute_insights(csv)["highlights"]]
    assert "You moved Rp 1.500.000 to your own savings or e-wallets. It counts as savings, not spending." in texts
    texts = [h["text"] for h in compute_insights(csv, lang="id")["highlights"]]
    assert any(x.startswith("Anda memindahkan Rp 1.500.000 ke tabungan") for x in texts)


# ---- saved = left over + moved to your savings ------------------------------


def saved_month():
    """Money in 10.000.000; total out 9.050.000, of which 1.500.000 to own savings."""
    rows = [
        tx("KR OTOMATIS GAJI PT CONTOH", 10_000_000, "other_transfer", "CR", "2026-09-01"),
        tx("TRSF E-BANKING DB PAK KOS CONTOH", 3_000_000, "housing", date="2026-09-02"),
        tx("INDOMARET CONTOH", 1_250_000, "groceries", date="2026-09-03"),
        tx("WARUNG CONTOH", 1_800_000, "eating_out", date="2026-09-04"),
        tx("TRSF KE IBU CONTOH", 1_500_000, "people_transfer", date="2026-09-05"),
        tx("TOPUP DANA CONTOH", 1_500_000, "savings_transfer", date="2026-09-30"),
    ]
    for i, r in enumerate(rows):
        r["id"] = i
    return rows


def test_saved_by_hand_is_the_same_everywhere():
    ins = compute_insights(saved_month())
    cash, b = ins["cash_flow"], ins["budget_503020"]
    assert cash["total_out"] == 9_050_000
    assert cash["left_over"] == 950_000  # 10.000.000 - 9.050.000
    assert cash["moved_to_savings"] == 1_500_000
    assert cash["saved"] == 2_450_000  # 950.000 + 1.500.000
    assert cash["savings_rate"] == 0.245
    # The 50/30/20 card says the same (one month, income from money in).
    assert b["savings"] == b["actual_amounts"]["savings"] == 2_450_000
    assert b["left_over"] == b["left_over_monthly"] == 950_000
    assert b["savings_transfers"] == b["savings_transfers_monthly"] == 1_500_000
    assert b["shares"]["savings"] == 0.245
    assert b["difference"]["savings"] == 2_000_000 - 2_450_000


def test_saved_highlight_in_english_and_indonesian():
    en = [h["text"] for h in compute_insights(saved_month())["highlights"]]
    assert "You saved Rp 2.450.000 (24.5%): Rp 950.000 left over + Rp 1.500.000 moved to your savings." in en
    id_ = [h["text"] for h in compute_insights(saved_month(), lang="id")["highlights"]]
    assert "Anda menabung Rp 2.450.000 (24,5%): sisa Rp 950.000 + Rp 1.500.000 dipindahkan ke tabungan Anda." in id_
    # Nothing moved: no "Rp 0 moved" breakdown.
    plain = [r for r in saved_month() if r["predicted"] != "savings_transfer"]
    texts = [h["text"] for h in compute_insights(plain)["highlights"]]
    assert "You saved Rp 2.450.000 (24.5%) of the Rp 10.000.000 that came in, all of it left over." in texts
    # Spending over money in.
    over = plain + [tx("TOKO CONTOH", 6_000_000, "shopping", date="2026-09-15")]
    texts = [h["text"] for h in compute_insights(over)["highlights"]]
    assert "Your spending was Rp 3.550.000 more than the Rp 10.000.000 that came in this period." in texts


def test_saved_without_money_in_is_none():
    csv = [dict(r, direction=None) for r in saved_month() if r["direction"] != "CR"]
    cash = compute_insights(csv)["cash_flow"]
    assert (cash["left_over"], cash["saved"], cash["savings_rate"]) == (None, None, None)
    assert cash["moved_to_savings"] == 1_500_000
    b = compute_insights(csv, income_override=10_000_000)["budget_503020"]
    assert b["actual_amounts"]["savings"] == 2_450_000 and b["left_over_monthly"] == 950_000


def test_reports_show_the_saved_breakdown(client):
    rows = saved_month()
    pdf = client.post("/api/report", json={"rows": rows}, params={"format": "pdf"})
    with pdfplumber.open(io.BytesIO(pdf.content)) as doc:
        text = " ".join((page.extract_text() or "").replace(chr(10), " ") for page in doc.pages)
    assert "Saved per month: Rp 2.450.000 = Rp 950.000 left over + Rp 1.500.000 moved to your savings." in text
    xlsx = client.post("/api/report", json={"rows": rows}, params={"format": "xlsx"})
    cells = {r[0]: r[1:] for r in load_workbook(io.BytesIO(xlsx.content))["Insights"].iter_rows(values_only=True)
             if r and r[0] is not None}
    assert cells["Left over (IDR)"][0] == 950_000
    assert cells["Moved to your savings (IDR)"][0] == 1_500_000
    assert cells["Saved (IDR)"][0] == 2_450_000 and cells["Savings rate"][0] == 0.245
    assert "Saved per month: Rp 2.450.000 = Rp 950.000 left over + Rp 1.500.000 moved to your savings." in cells
    pdf_id = client.post("/api/report", json={"rows": rows}, params={"format": "pdf", "lang": "id"})
    with pdfplumber.open(io.BytesIO(pdf_id.content)) as doc:
        text = " ".join((page.extract_text() or "").replace(chr(10), " ") for page in doc.pages)
    assert "Ditabung per bulan: Rp 2.450.000 = sisa Rp 950.000 + Rp 1.500.000 dipindahkan ke tabungan Anda." in text


def test_top_transfers_cover_the_transfers_group_and_each_category():
    top = compute_insights(month())["top_transfers"]
    # savings_transfer is a transfer too: its recipient is listed with the others.
    assert [m["name"] for m in top["out_by_total"]] == ["Topup Dana Contoh", "Ibu Contoh", "Cash withdrawal (ATM)"]
    assert set(top["by_category"]) == {"savings_transfer", "people_transfer", "cash", "other_transfer"}
    people = top["by_category"]["people_transfer"]
    assert people["out_by_total"] == [{"name": "Ibu Contoh", "count": 1, "total": 700_000, "avg": 700_000, "share": 1.0}]
    assert people["in_by_total"] == []
    assert top["by_category"]["other_transfer"]["in_by_total"][0]["name"] == "Gaji PT Contoh"
    assert top["by_category"]["other_transfer"]["out_by_total"] == []
    # Merchants never include the Transfers group.
    assert all(m["merchant"] != "Topup Dana Contoh" for m in compute_insights(month())["top_merchants"]["by_total"])


# ---- reports ---------------------------------------------------------------


def test_reports_have_group_subtotals_and_other_money_out(client):
    rows = month()
    pdf = client.post("/api/report", json={"rows": rows}, params={"format": "pdf"})
    with pdfplumber.open(io.BytesIO(pdf.content)) as doc:
        text = "\n".join(page.extract_text() or "" for page in doc.pages)
    assert "Needs subtotal 3 Rp 3.500.000" in text
    assert "Wants subtotal 2 Rp 1.000.000" in text
    assert "Transfers subtotal 3 Rp 2.500.000" in text
    assert "Other money out (family, cash, other" in text and "Rp 1.000.000" in text

    xlsx = client.post("/api/report", json={"rows": rows}, params={"format": "xlsx"})
    wb = load_workbook(io.BytesIO(xlsx.content))
    summary = {r[0]: r[1:] for r in wb["Summary"].iter_rows(values_only=True) if r and r[0] is not None}
    assert "Totals by group" in summary
    assert summary["Needs subtotal"][:2] == (3, 3_500_000)
    assert summary["Transfers subtotal"][:2] == (3, 2_500_000)
    assert summary["Wants subtotal"][2] == pytest.approx(1_000_000 / 7_000_000, abs=1e-4)
    insights = {r[0]: r[1:] for r in wb["Insights"].iter_rows(values_only=True) if r and r[0] is not None}
    assert insights["Other money out (family, cash, other transfers)"][1] == 1_000_000
    assert insights["Savings (left over + moved to savings)"][:3] == (2_000_000, 4_500_000, -2_500_000)

    xlsx_id = client.post("/api/report", json={"rows": rows}, params={"format": "xlsx", "lang": "id"})
    labels = [r[0] for r in load_workbook(io.BytesIO(xlsx_id.content))["Ringkasan"].iter_rows(values_only=True)]
    assert "Subtotal Kebutuhan" in labels and "Total per kelompok" in labels


def test_report_source_memory_label(client):
    rows = [dict(r, source="memory") for r in month()[:2]]
    pdf = client.post("/api/report", json={"rows": rows}, params={"format": "pdf"})
    assert pdf.status_code == 200
    with pdfplumber.open(io.BytesIO(pdf.content)) as doc:
        assert "memory" in "\n".join(page.extract_text() or "" for page in doc.pages)


def test_clean_description_unchanged_for_memory_keys():
    # The memory key is computed on the raw description, not the cleaned one.
    assert clean_description("TRSF E-BANKING DB 0109/FTSCY/WS1 IBU CONTOH") == "transfer e banking ibu contoh"
