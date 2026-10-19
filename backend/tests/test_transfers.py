"""Top transfers: the counterparty key, grouping and shares, the recipient highlight, the report
tables and privacy. Synthetic rows only: every person and company below is made up.
"""

import io
import logging

import pdfplumber
import pytest
from openpyxl import load_workbook

from classifier.insights import TRANSFER_CASH, TRANSFER_OTHER, compute_insights, transfer_key
from test_api import post_csv
from test_insights import client, statement, tx  # noqa: F401 (client is a fixture)


@pytest.mark.parametrize("description, key", [
    ("TRSF E-BANKING DB 0309/FTSCY/WS95051 BUDI CONTOH", "Budi Contoh"),
    ("TRSF KE BCA BUDI SANTOSO", "Budi Santoso"),
    ("SWITCHING DB TRANSFER KE 014 SITI CONTOH", "Siti Contoh"),
    ("KR OTOMATIS LLG-ANTAR BANK DARI ANDI CONTOH", "Andi Contoh"),
    ("TRSF E-BANKING CR 2509/ACCSW/WS95031 PT CONTOH JAYA ABADI SEJAHTERA", "PT Contoh Jaya Abadi"),
    ("BI-FAST DB TRANSFER KE 09/12 RINA CONTOH", "Rina Contoh"),
    ("BI-FAST TRANSFER REF 91", TRANSFER_OTHER),
    ("TRSF E-BANKING DB 0309/FTSCY/WS95051", TRANSFER_OTHER),
    ("", TRANSFER_OTHER),
    ("TARIK TUNAI ATM 12/09 CABANG CONTOH", TRANSFER_CASH),
    ("ATM BCA 0012", TRANSFER_CASH),
])
def test_transfer_key(description, key):
    assert transfer_key(description) == key


def transfers_statement():
    """Salary in, rent and family transfers out, an unnamed transfer, a cash withdrawal and food."""
    rows = [
        tx("KR OTOMATIS LLG-ANTAR BANK DARI PT CONTOH GAJI", 10_000_000, "transfer", "CR", "2026-09-01"),
        tx("TRSF E-BANKING CR 0509/FTSCY/WS1 IBU CONTOH", 500_000, "transfer", "CR", "2026-09-05"),
        tx("TRSF E-BANKING DB 0209/FTSCY/WS95051 PAK KOS CONTOH", 2_500_000, "transfer", date="2026-09-02"),
        tx("TRSF E-BANKING DB 0309/FTSCY/WS95052 IBU CONTOH", 1_000_000, "transfer", date="2026-09-03"),
        tx("TRSF E-BANKING DB 2009/FTSCY/WS95053 IBU CONTOH", 1_000_000, "transfer", date="2026-09-20"),
        tx("BI-FAST TRANSFER REF 91", 300_000, "transfer", date="2026-09-10"),
        tx("TARIK TUNAI ATM 12/09", 200_000, "transfer", date="2026-09-12"),
        tx("GRAB FOOD 1122", 100_000, "food", date="2026-09-13"),
    ]
    for i, r in enumerate(rows):
        r["id"] = i
    return rows


def test_top_transfers_group_out_and_in():
    top = compute_insights(transfers_statement())["top_transfers"]
    out_total = 2_500_000 + 1_000_000 + 1_000_000 + 300_000 + 200_000
    assert [m["name"] for m in top["out_by_total"]] == [
        "Pak Kos Contoh", "Ibu Contoh", "Other transfer", "Cash withdrawal (ATM)"]
    first, second = top["out_by_total"][:2]
    assert first == {"name": "Pak Kos Contoh", "count": 1, "total": 2_500_000, "avg": 2_500_000,
                     "share": round(2_500_000 / out_total, 4)}
    assert second["count"] == 2 and second["total"] == 2_000_000 and second["avg"] == 1_000_000
    assert top["out_by_count"][0]["name"] == "Ibu Contoh"
    # Money in: share of all money in; non-transfer rows never appear.
    assert [(m["name"], m["share"]) for m in top["in_by_total"]] == [
        ("PT Contoh Gaji", round(10_000_000 / 10_500_000, 4)), ("Ibu Contoh", round(500_000 / 10_500_000, 4))]
    assert all("Grab" not in m["name"] for m in top["out_by_total"] + top["out_by_count"])


def test_top_transfers_names_are_translated():
    top = compute_insights(transfers_statement(), lang="id")["top_transfers"]
    names = [m["name"] for m in top["out_by_total"]]
    assert "Transfer lainnya" in names and "Tarik tunai (ATM)" in names


def test_top_transfers_empty_without_transfers():
    rows = [tx("GRAB FOOD 1122", 100_000, "food")]
    assert compute_insights(rows)["top_transfers"] == {"out_by_total": [], "out_by_count": [], "in_by_total": []}


def test_top_transfers_without_amounts_have_no_share():
    rows = [tx("TRSF KE BUDI CONTOH", None, "transfer")]
    top = compute_insights(rows)["top_transfers"]
    assert top["out_by_total"] == []  # nothing to rank by total
    assert top["out_by_count"] == [{"name": "Budi Contoh", "count": 1, "total": 0, "avg": 0, "share": None}]


def test_recipient_highlight_in_both_languages():
    hl_en = [h["text"] for h in compute_insights(transfers_statement())["highlights"]]
    assert any(h.startswith("Your largest transfer recipient was Pak Kos Contoh: Rp 2.500.000 over 1 payment")
               for h in hl_en), hl_en
    hl_id = [h["text"] for h in compute_insights(transfers_statement(), lang="id")["highlights"]]
    assert any(h.startswith("Penerima transfer terbesar Anda adalah Pak Kos Contoh") for h in hl_id), hl_id


def test_recipient_highlight_skips_unnamed_and_small_recipients():
    # The largest group is unnamed, so the highlight names the largest named recipient instead.
    rows = [
        tx("BI-FAST TRANSFER REF 91", 5_000_000, "transfer"),
        tx("TRSF KE BUDI CONTOH", 1_200_000, "transfer"),
    ]
    texts = [h["text"] for h in compute_insights(rows)["highlights"]]
    assert any("recipient was Budi Contoh" in h for h in texts)
    assert not any("Other transfer" in h for h in texts)

    # Under 20% of transfers out and under 10% of all money out: no recipient highlight.
    names = ["SATU", "DUA", "TIGA", "EMPAT", "LIMA", "ENAM", "TUJUH", "DELAPAN", "SEMBILAN", "SEPULUH"]
    rows = [tx(f"TRSF KE ORANG {name}", 100_000, "transfer") for name in names]
    rows.append(tx("TOKO CONTOH", 20_000_000, "shopping"))
    assert not any("recipient" in h["text"] for h in compute_insights(rows)["highlights"])


def test_insights_endpoint_returns_top_transfers(client):
    body = client.post("/api/insights", json={"rows": transfers_statement()}).json()
    assert body["top_transfers"]["out_by_total"][0]["name"] == "Pak Kos Contoh"
    assert set(body["top_transfers"]) == {"out_by_total", "out_by_count", "in_by_total"}


def test_pdf_has_top_transfers_table(client):
    resp = client.post("/api/report", json={"rows": transfers_statement()}, params={"format": "pdf"})
    assert resp.status_code == 200
    with pdfplumber.open(io.BytesIO(resp.content)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "Top transfers (out)" in text and "Pak Kos Contoh" in text and "Cash withdrawal (ATM)" in text

    resp = client.post("/api/report", json={"rows": transfers_statement()}, params={"format": "pdf", "lang": "id"})
    with pdfplumber.open(io.BytesIO(resp.content)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "Transfer teratas (keluar)" in text and "Tarik tunai (ATM)" in text


def test_xlsx_has_top_transfers_block(client):
    resp = client.post("/api/report", json={"rows": transfers_statement()}, params={"format": "xlsx"})
    ws = load_workbook(io.BytesIO(resp.content))["Insights"]
    values = [[c.value for c in row] for row in ws.iter_rows()]
    start = next(i for i, row in enumerate(values) if row[0] == "Top transfers (out)")
    assert values[start + 1][:5] == ["Recipient", "Count", "Total (IDR)", "Average (IDR)", "Share"]
    assert values[start + 2][:4] == ["Pak Kos Contoh", 1, 2_500_000, 2_500_000]
    assert ws.cell(row=start + 3, column=5).number_format == "0.0%"


def test_logs_hold_no_recipient_names(client, caplog):
    caplog.set_level(logging.DEBUG)
    rows = transfers_statement()
    assert client.post("/api/insights", json={"rows": rows}).status_code == 200
    for fmt in ("csv", "xlsx", "pdf"):
        assert client.post("/api/report", json={"rows": rows}, params={"format": fmt}).status_code == 200
    csv = "date,description,amount\n" + "\n".join(
        f"{r['date']},{r['description']},{r['amount']}" for r in rows if r["direction"] == "DB")
    assert post_csv(client, csv.encode()).status_code == 200

    text = "\n".join(r.getMessage() for r in caplog.records)
    assert "POST /api/insights 200" in text and "POST /api/report 200" in text and "POST /api/classify 200" in text
    for secret in ["Pak Kos", "PAK KOS", "Ibu Contoh", "IBU CONTOH", "Contoh Gaji", "2500000", "2.500.000"]:
        assert secret not in text


def test_statement_fixture_still_has_one_recipient():
    # The shared insights fixture: one named transfer out, 100% of transfers.
    top = compute_insights(statement())["top_transfers"]
    assert top["out_by_total"] == [{"name": "Budi Contoh", "count": 1, "total": 2_000_000, "avg": 2_000_000,
                                    "share": 1.0}]
