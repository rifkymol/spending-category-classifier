import csv
import io
import json
import logging
import re

import pdfplumber
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm

from app.main import CATEGORIES, MAX_REPORT_BYTES, app
from app.reports import BUDGET_COL_MM, CATEGORY_COLORS, OTHER_COLOR, _budget_table, format_idr, period, safe_text
from app.schemas import MAX_REPORT_ROWS, Row
from classifier.insights import compute_insights
from test_api import ROOT, SAMPLE, post_csv, project_files

INJECTION = '=HYPERLINK("http://x")'
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


@pytest.fixture(scope="module")
def rows(client):
    rows = post_csv(client, SAMPLE.read_bytes()).json()["rows"]
    rows[0]["description"] = INJECTION
    rows[1]["amount"] = -45000.0
    return rows


def get_report(client, rows, fmt=None):
    params = {"format": fmt} if fmt else {}
    return client.post("/api/report", json={"rows": rows}, params=params)


def expected_totals(rows) -> dict[str, float]:
    totals = {c: 0.0 for c in CATEGORIES}
    for r in rows:
        totals[r["predicted"]] += abs(r["amount"] or 0)
    return totals


# ---- helpers ---------------------------------------------------------------


@pytest.mark.parametrize("value", ["=1+1", "+62", "-5", "@SUM(A1)", "\tx", "\rx"])
def test_safe_text_prefixes_formula_starts(value):
    assert safe_text(value) == "'" + value


def test_safe_text_leaves_plain_text():
    assert safe_text("GRAB*FOOD") == "GRAB*FOOD"
    assert safe_text("") == ""


def test_format_idr():
    assert format_idr(6731990) == "Rp 6.731.990"
    assert format_idr(-45000) == "-Rp 45.000"
    assert format_idr(0) == "Rp 0"


def _row(date):
    return Row(id=0, date=date, description="x", amount=1.0, predicted="eating_out", confidence=1.0,
               needs_review=False, source="rule")


def test_period():
    assert period([_row("2026-09-01"), _row("2026-09-21"), _row("2026-09-05")]) == "2026-09-01 to 2026-09-21"
    assert period([_row("02/09/2026"), _row("18/09/2026")]) == "2026-09-02 to 2026-09-18"
    assert period([_row(None), _row("not a date")]) is None


def test_pdf_colors_match_frontend():
    # Both sides come from classifier/categories.py: the reports directly, the frontend through the
    # exported categories.json (test_categories.py checks that file is up to date).
    data = json.loads((ROOT / "frontend" / "src" / "lib" / "categories.json").read_text(encoding="utf-8"))
    assert {c["key"]: c["color"] for c in data["categories"]} == CATEGORY_COLORS
    assert data["other_color"] == OTHER_COLOR
    assert list(CATEGORY_COLORS) == CATEGORIES


# ---- formats ---------------------------------------------------------------


def test_csv_is_default_and_neutralizes_injection(client, rows):
    resp = get_report(client, rows)
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    assert resp.headers["content-disposition"] == 'attachment; filename="spending-report.csv"'

    table = list(csv.DictReader(io.StringIO(resp.text.split("\n\n", 1)[1])))
    assert table[0]["description"] == "'" + INJECTION
    assert table[1]["amount"] == "-45000.0"  # numbers are never prefixed
    assert get_report(client, rows, "csv").text == resp.text


def test_xlsx_report(client, rows):
    resp = get_report(client, rows, "xlsx")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == XLSX_TYPE
    assert resp.headers["content-disposition"] == 'attachment; filename="spending-report.xlsx"'

    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Summary", "Transactions", "Insights"]

    summary = wb["Summary"]
    assert summary._charts, "the Summary sheet has a pie chart"
    table = {}
    for r in summary.iter_rows(values_only=True):
        if r[0] in CATEGORIES:
            table[r[0]] = r
    want = expected_totals(rows)
    assert set(table) == set(CATEGORIES)
    for cat, (_, count, total, share) in table.items():
        assert count == sum(r["predicted"] == cat for r in rows)
        assert total == pytest.approx(want[cat])
    assert sum(r[3] for r in table.values()) == pytest.approx(1.0, abs=1e-3)

    tx = wb["Transactions"]
    assert tx.freeze_panes == "A2"
    assert tx.auto_filter.ref
    header = [c.value for c in tx[1]]
    desc_col = header.index("description") + 1
    amount_col = header.index("amount") + 1
    injected = tx.cell(row=2, column=desc_col)
    assert injected.data_type != "f"
    assert injected.value == "'" + INJECTION
    negative = tx.cell(row=3, column=amount_col)
    assert negative.value == -45000 and negative.data_type == "n"
    assert tx.max_row == len(rows) + 1
    assert not any(c.data_type == "f" for row in tx.iter_rows() for c in row)


def test_pdf_report(client, rows):
    resp = get_report(client, rows, "pdf")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert resp.headers["content-disposition"] == 'attachment; filename="spending-report.pdf"'
    assert resp.content.startswith(b"%PDF")
    assert len(resp.content) > 1024
    # Built-in Helvetica only; reportlab drawings default to Times-Roman unless told otherwise.
    fonts = set(re.findall(rb"/BaseFont\s*/([\w-]+)", resp.content))
    assert fonts and all(f.startswith(b"Helvetica") for f in fonts), fonts
    assert b"/Times-Roman" not in resp.content


def pdf_pages(content: bytes) -> list[str]:
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        return [page.extract_text() or "" for page in pdf.pages]


def test_pdf_budget_table_fits_a4(client):
    sample = post_csv(client, SAMPLE.read_bytes()).json()["rows"]
    budget = compute_insights([Row(**r) for r in sample])["budget_503020"]
    cell = getSampleStyleSheet()["BodyText"].clone("cell", fontSize=8, leading=10)
    income, _, table, _, caption = _budget_table(budget, cell)
    assert table._colWidths == [w * mm for w in BUDGET_COL_MM]
    assert sum(table._colWidths) <= 180 * mm  # A4 width minus 15 mm margins
    assert [p.text for p in table._cellvalues[0]] == [
        "50/30/20 view", "Target", "Actual", "Difference", "Share", "Guideline"]
    assert table.repeatRows == 1
    assert income.text.startswith("Income (per month): not available")
    assert caption.text == "Actual is per month. Share = share of total out (no income found)."


def test_pdf_headers_do_not_run_together(client, rows):
    text = "\n".join(pdf_pages(get_report(client, rows, "pdf").content))
    assert "50/30/20 view Target Actual Difference Share Guideline" in text
    assert "TargetActual" not in text
    assert "Actual is per month. Share = share of total out (no income found)." in text


def test_reports_with_income_override_have_target_columns(client, rows):
    """A CSV upload has no money in; an entered income gives the PDF and Excel Rupiah targets."""
    body = {"rows": rows, "income_override": 10_000_000}
    pdf = client.post("/api/report", json=body, params={"format": "pdf"})
    assert pdf.status_code == 200
    text = "\n".join(pdf_pages(pdf.content))
    assert "Income (per month): Rp 10.000.000 (entered)" in text
    assert "Rp 5.000.000" in text and "Rp 3.000.000" in text and "Rp 2.000.000" in text  # 50/30/20 targets
    assert re.search(r"Rp [\d.]+ (left|over)", text)

    xlsx = client.post("/api/report", json=body, params={"format": "xlsx"})
    ws = load_workbook(io.BytesIO(xlsx.content))["Insights"]
    cells = {r[0]: r[1:] for r in ws.iter_rows(values_only=True) if r and r[0] is not None}
    assert cells["Bucket"][:5] == ("Target (IDR)", "Actual (IDR)", "Difference (IDR)", "Share", "Guideline")
    budget = compute_insights([Row(**r) for r in rows], 10_000_000)["budget_503020"]
    needs = cells["Needs (essentials)"]
    assert needs[0] == 5_000_000
    assert needs[1] == pytest.approx(budget["actual_amounts"]["needs"])
    assert needs[2] == pytest.approx(5_000_000 - budget["actual_amounts"]["needs"])
    assert "Income (per month): Rp 10.000.000 (entered)" in cells


def test_reports_in_indonesian(client, rows):
    body = {"rows": rows, "income_override": 10_000_000}
    pdf = client.post("/api/report", json=body, params={"format": "pdf", "lang": "id"})
    text = "\n".join(pdf_pages(pdf.content))
    for phrase in ["Laporan pengeluaran", "Sorotan", "Tampilan 50/30/20", "Target Aktual Selisih",
                   "Pendapatan (per bulan): Rp 10.000.000 (diisi manual)", "Wawasan anggaran, bukan nasihat keuangan.",
                   "Makan di luar", "Tanggal Keterangan Nominal Kategori Sumber", "Halaman 1"]:
        assert phrase in text, phrase
    assert "Spending report" not in text and "Highlights" not in text

    xlsx = client.post("/api/report", json=body, params={"format": "xlsx", "lang": "id"})
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == ["Ringkasan", "Transaksi", "Wawasan"]
    cells = {r[0]: r[1:] for r in wb["Wawasan"].iter_rows(values_only=True) if r and r[0] is not None}
    assert "Kebutuhan (pokok)" in cells and "Arus kas" in cells
    assert "Wawasan anggaran, bukan nasihat keuangan." in cells
    summary_labels = [r[0] for r in wb["Ringkasan"].iter_rows(values_only=True)]
    assert "Makan di luar" in summary_labels and "Total pengeluaran (IDR)" in summary_labels

    # The CSV report is a machine format: English headers and category keys in both languages.
    csv_en = client.post("/api/report", json=body, params={"format": "csv"}).text
    assert client.post("/api/report", json=body, params={"format": "csv", "lang": "id"}).text == csv_en


def test_invalid_lang_falls_back_to_english(client, rows):
    resp = client.post("/api/report", json={"rows": rows}, params={"format": "pdf", "lang": "fr"})
    assert resp.status_code == 200
    assert "Spending report" in pdf_pages(resp.content)[0]


def test_pdf_transaction_header_repeats_on_every_page(client, rows):
    many = [dict(r, id=i) for i, r in enumerate(rows * 6)]
    pages = pdf_pages(get_report(client, many, "pdf").content)
    header = "Date Description Amount Category Source"
    first = next(i for i, p in enumerate(pages) if header in p)
    assert len(pages) - first >= 2, "the transactions table should span several pages"
    assert all(header in p for p in pages[first:])
    # No table header is left alone at the bottom of a page: each page that shows a header row
    # also shows at least one data row under it.
    for p in pages:
        lines = p.splitlines()
        for i, line in enumerate(lines):
            if line.startswith(("Category Count Total Share", "Top merchants by total", header)):
                assert i + 2 < len(lines), f"orphaned header: {line!r}"


def test_reports_handle_odd_text_and_missing_values(client, rows):
    odd = [dict(r) for r in rows[:3]]
    odd[0]["description"] = "<b>CAFÉ & co</b> \x01 ☕"
    odd[1]["date"] = None
    odd[1]["amount"] = None
    for fmt in ("csv", "xlsx", "pdf"):
        assert get_report(client, odd, fmt).status_code == 200, fmt


def test_reports_handle_description_only_rows(client):
    rows = post_csv(client, b"description\nGRAB*FOOD 123\nSOMETHING UNKNOWN\n").json()["rows"]
    for fmt in ("csv", "xlsx", "pdf"):
        assert get_report(client, rows, fmt).status_code == 200, fmt


# ---- validation ------------------------------------------------------------


def test_unknown_format_is_422(client, rows):
    assert get_report(client, rows, "docx").status_code == 422


def test_bad_category_is_422_without_echoing_row_data(client, rows):
    bad = [dict(r) for r in rows[:2]]
    bad[0]["predicted"] = "groceries-SECRET-CAT"
    bad[0]["description"] = "PRIVATE MERCHANT XYZ"
    resp = get_report(client, bad, "xlsx")
    assert resp.status_code == 422
    detail = resp.json()["detail"]
    assert detail[0]["loc"] == ["body", "rows", 0, "predicted"]
    assert all(set(e) == {"loc", "msg", "type"} for e in detail)
    for secret in ["PRIVATE MERCHANT", "groceries-SECRET-CAT", "input"]:
        assert secret not in resp.text


def test_malformed_json_is_422_without_echo(client):
    resp = client.post("/api/report", content=b'{"rows": [{"description": "PRIVATE MERCHANT"', headers={"Content-Type": "application/json"})
    assert resp.status_code == 422
    assert "PRIVATE MERCHANT" not in resp.text


def test_too_many_rows_is_rejected(client):
    row = {"id": 0, "date": None, "description": "x", "amount": None, "predicted": "food",
           "confidence": 1.0, "needs_review": False, "source": "rule"}
    body = json.dumps({"rows": [row] * (MAX_REPORT_ROWS + 1)}).encode()
    assert len(body) < MAX_REPORT_BYTES  # so the row cap, not the size cap, is what rejects it
    resp = client.post("/api/report", content=body, headers={"Content-Type": "application/json"})
    assert resp.status_code in (413, 422)
    assert "50000" in resp.text


def test_body_over_10mb_is_413(client):
    body = b'{"rows": [' + b" " * MAX_REPORT_BYTES + b"]}"
    resp = client.post("/api/report", content=body, headers={"Content-Type": "application/json"})
    assert resp.status_code == 413


# ---- privacy ---------------------------------------------------------------


def test_reports_write_no_files(client, rows, isolated_tmp):
    before_project = project_files()
    for fmt in ("csv", "xlsx", "pdf"):
        assert get_report(client, rows, fmt).status_code == 200
    assert project_files() - before_project == set()
    assert list(isolated_tmp.iterdir()) == []


def test_report_logs_hold_no_descriptions_or_amounts(client, rows, caplog):
    caplog.set_level(logging.DEBUG)
    for fmt in ("csv", "xlsx", "pdf"):
        get_report(client, rows, fmt)
    bad = [dict(rows[2], predicted="nope")]
    get_report(client, bad, "pdf")

    # Formatted messages only: raw args hold the duration float, whose digits can look like amounts.
    text = "\n".join(r.getMessage() for r in caplog.records)
    assert "POST /api/report 200" in text and f"rows={len(rows)}" in text  # the logger did run
    assert "POST /api/report 422" in text
    secrets = ["HYPERLINK", "45000", rows[2]["description"], rows[3]["description"]]
    secrets += [str(int(r["amount"])) for r in rows[2:6] if r["amount"]]
    for secret in secrets:
        assert secret not in text
