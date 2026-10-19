"""Insights: metrics, edge cases, the /api/insights endpoint, privacy and report sections.

Synthetic rows only: every merchant and person below is made up.
"""

import io
import json
import logging

import pdfplumber
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import MAX_REPORT_BYTES, app
from app.reports import report_xlsx
from app.schemas import Row
from classifier.insights import (
    DISCLAIMER,
    SMALL_SPEND_THRESHOLD,
    compute_insights,
    format_rp,
    is_fee,
    merchant_key,
)
from test_api import SAMPLE, post_csv, project_files

INJECTION = '=HYPERLINK("http://x")'


def tx(description, amount, predicted, direction="DB", date="2026-09-01", **extra):
    return {"id": extra.get("id", 0), "date": date, "description": description, "amount": amount,
            "direction": direction, "predicted": predicted, "confidence": 1.0, "needs_review": False,
            "source": "rule"}


def statement():
    """One synthetic month: salary in, living spend, a transfer, a fee and small spends."""
    rows = [
        tx("KR OTOMATIS GAJI PT CONTOH", 10_000_000, "transfer", "CR", "2026-09-01"),
        tx("QR 00012345 WARUNG MAJU JAYA 998877", 30_000, "food", date="2026-09-02"),
        tx("QR 00067890 WARUNG MAJU JAYA 112233", 25_000, "food", date="2026-09-03"),
        tx("TRANSAKSI DEBIT TGL: 04/09 TOKO SERBA ADA 5544", 1_500_000, "shopping", date="2026-09-04"),
        tx("PLN PRABAYAR 1234567890", 500_000, "bills", date="2026-09-05"),
        tx("GOJEK RIDE 4411", 20_000, "transport", date="2026-09-06"),
        tx("TRSF E-BANKING DB 0309/FTSCY/WS95051 BUDI CONTOH", 2_000_000, "transfer", date="2026-09-07"),
        tx("BIAYA ADM", 15_000, "bills", date="2026-09-30"),
    ]
    for i, r in enumerate(rows):
        r["id"] = i
    return rows


# ---- metrics ---------------------------------------------------------------


def test_cash_flow():
    cash = compute_insights(statement())["cash_flow"]
    out = 30_000 + 25_000 + 1_500_000 + 500_000 + 20_000 + 2_000_000 + 15_000
    assert cash["money_in"] == 10_000_000 and cash["money_in_count"] == 1
    assert cash["total_out"] == out and cash["out_count"] == 7
    assert cash["transfers_out"] == 2_000_000
    assert cash["living_spend"] == out - 2_000_000
    assert cash["net"] == 10_000_000 - out
    assert cash["savings_rate"] == pytest.approx((10_000_000 - out) / 10_000_000, abs=1e-4)


def test_top_merchants_group_and_exclude_transfers():
    top = compute_insights(statement())["top_merchants"]
    names = [m["merchant"] for m in top["by_total"]]
    assert names[0] == "Toko Serba Ada"
    assert "Budi Contoh" not in " ".join(names) and all("Budi" not in n for n in names)
    warung = next(m for m in top["by_count"] if m["merchant"] == "Warung Maju Jaya")
    assert warung == {"merchant": "Warung Maju Jaya", "count": 2, "total": 55_000, "avg": 27_500}
    assert top["by_count"][0]["merchant"] == "Warung Maju Jaya"
    assert len(top["by_total"]) <= 5 and len(top["by_count"]) <= 5


@pytest.mark.parametrize("description, key", [
    ("QR 00012345 WARUNG MAJU JAYA 998877", "Warung Maju Jaya"),
    ("QRIS KOPI CONTOH JKT 4411", "Kopi Contoh Jkt"),
    ("TRANSAKSI DEBIT TGL: 01/09 QR 014 BAKSO FIKTIF RASA NIKMAT", "Bakso Fiktif Rasa"),
    ("GRAB*FOOD 2309011234", "Grab Food"),
    ("TRSF E-BANKING DB 0309/FTSCY/WS95051", ""),
])
def test_merchant_key(description, key):
    assert merchant_key(description) == key


def test_small_spend_boundary():
    rows = [tx("TOKO A", SMALL_SPEND_THRESHOLD - 1, "shopping"), tx("TOKO B", SMALL_SPEND_THRESHOLD, "shopping"),
            tx("TEMAN", 10_000, "transfer")]
    small = compute_insights(rows)["small_spends"]
    assert small["threshold"] == 50_000
    assert small["count"] == 1 and small["total"] == 49_999
    assert small["share_of_living"] == pytest.approx(49_999 / 99_999, abs=1e-4)


def test_fees():
    assert is_fee("BIAYA ADM") and is_fee("Biaya Administrasi") and is_fee("BIAYA TRANSFER BI-FAST")
    assert is_fee("ADMIN FEE 2026") and not is_fee("WARUNG MAJU JAYA")
    fees = compute_insights(statement())["fees"]
    assert fees == {"count": 1, "total": 15_000}


def test_budget_503020_uses_money_in():
    b = compute_insights(statement())["budget_503020"]
    assert b["basis"] == "money_in" and b["basis_amount"] == 10_000_000
    assert b["needs"] == 500_000 + 20_000 + 15_000
    assert b["wants"] == 30_000 + 25_000 + 1_500_000
    assert b["unclassified_transfers"] == 2_000_000
    assert b["savings"] == compute_insights(statement())["cash_flow"]["net"]
    assert b["shares"]["needs"] == pytest.approx(0.0535)
    assert b["shares"]["unclassified_transfers"] == pytest.approx(0.2)
    assert b["mapping"] == {"needs": ["bills", "transport"], "wants": ["food", "shopping"], "unclassified": ["transfer"]}


def test_period_and_per_month():
    ins = compute_insights(statement())
    assert ins["period"] == {"start": "2026-09-01", "end": "2026-09-30", "days": 30, "months_factor": 1.0}
    assert ins["per_month"]["shopping"] == 1_500_000

    two_months = [tx("TOKO A", 100_000, "shopping", date="2026-08-01"),
                  tx("TOKO A", 100_000, "shopping", date="2026-09-30")]
    ins = compute_insights(two_months)
    assert ins["period"]["days"] == 61
    assert ins["per_month"]["shopping"] == pytest.approx(200_000 * 30.4375 / 61, abs=0.01)


def test_highlights_sorted_and_counted():
    h = compute_insights(statement())["highlights"]
    assert 3 <= len(h) <= 5
    impacts = [x["impact_rp"] for x in h]
    assert impacts == sorted(impacts, reverse=True)
    assert all("Rp " in x["text"] for x in h)
    assert any("Rp 5.910.000" in x["text"] for x in h)  # kept: 10.000.000 - 4.090.000
    assert all("invest" not in x["text"].lower() for x in h)


def test_format_rp():
    assert format_rp(1_234_567) == "Rp 1.234.567"
    assert format_rp(-5_000) == "-Rp 5.000"


# ---- edge cases ------------------------------------------------------------


def test_no_money_in_uses_total_out():
    rows = [r for r in statement() if r["direction"] != "CR"]
    for r in rows:
        r["direction"] = None  # like a CSV upload
    ins = compute_insights(rows)
    assert ins["cash_flow"]["savings_rate"] is None
    b = ins["budget_503020"]
    assert b["basis"] == "total_out" and b["basis_amount"] == ins["cash_flow"]["total_out"]
    assert b["savings"] is None and b["shares"]["savings"] is None
    assert not any("came in" in h["text"] for h in ins["highlights"])


def test_no_amounts():
    rows = [tx("WARUNG MAJU JAYA", None, "food", date=None), tx("TOKO SERBA ADA", None, "shopping", date=None)]
    ins = compute_insights(rows)
    assert ins["cash_flow"]["total_out"] == 0 and ins["cash_flow"]["out_count"] == 2
    assert ins["small_spends"]["count"] == 0 and ins["small_spends"]["share_of_living"] is None
    assert ins["top_merchants"]["by_total"] == []
    assert [m["count"] for m in ins["top_merchants"]["by_count"]] == [1, 1]
    assert ins["period"] == {"start": None, "end": None, "days": None, "months_factor": 1.0}
    assert ins["highlights"] == []
    assert ins["budget_503020"]["shares"]["needs"] is None


def test_single_row():
    ins = compute_insights([tx("WARUNG MAJU JAYA", 30_000, "food")])
    assert ins["cash_flow"]["living_spend"] == 30_000
    assert ins["period"]["days"] == 1 and ins["period"]["months_factor"] == 1.0
    assert ins["top_merchants"]["by_total"][0]["merchant"] == "Warung Maju Jaya"
    assert ins["highlights"]


def test_all_transfers():
    rows = [tx("TRSF E-BANKING BUDI CONTOH", 500_000, "transfer"), tx("BI-FAST SITI CONTOH", 250_000, "transfer")]
    ins = compute_insights(rows)
    assert ins["cash_flow"]["living_spend"] == 0 and ins["cash_flow"]["transfers_out"] == 750_000
    assert ins["top_merchants"] == {"by_total": [], "by_count": []}
    assert ins["small_spends"]["share_of_living"] is None
    assert ins["budget_503020"]["shares"]["unclassified_transfers"] == 1.0
    assert len(ins["highlights"]) == 1 and "transfers" in ins["highlights"][0]["text"]


def test_empty_and_row_objects():
    assert compute_insights([])["highlights"] == []
    objs = [Row(**{k: v for k, v in r.items()}) for r in statement()]
    assert compute_insights(objs) == compute_insights(statement())


def test_json_serializable():
    ins = compute_insights(statement())
    assert json.loads(json.dumps(ins)) == ins
    assert ins["disclaimer"] == DISCLAIMER == "Budgeting insights, not financial advice."


# ---- API -------------------------------------------------------------------


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


def test_insights_endpoint(client):
    resp = client.post("/api/insights", json={"rows": statement()})
    assert resp.status_code == 200
    assert resp.json() == compute_insights(statement())


def test_insights_bad_category_is_422_without_echo(client):
    bad = statement()
    bad[1]["predicted"] = "groceries-SECRET-CAT"
    resp = client.post("/api/insights", json={"rows": bad})
    assert resp.status_code == 422
    detail = resp.json()["detail"]
    assert detail[0]["loc"] == ["body", "rows", 1, "predicted"]
    assert all(set(e) == {"loc", "msg", "type"} for e in detail)
    for secret in ["WARUNG", "groceries-SECRET-CAT", "input"]:
        assert secret not in resp.text


def test_insights_over_10mb_is_413(client):
    body = b'{"rows": [' + b" " * MAX_REPORT_BYTES + b"]}"
    resp = client.post("/api/insights", content=body, headers={"Content-Type": "application/json"})
    assert resp.status_code == 413


def test_classify_includes_insights(client):
    data = post_csv(client, SAMPLE.read_bytes()).json()
    assert {"rows", "summary", "insights"} <= set(data)
    assert data["insights"]["cash_flow"]["total_out"] == pytest.approx(data["summary"]["total_amount"])
    assert data["insights"]["disclaimer"] == DISCLAIMER


def test_insights_write_no_files(client, isolated_tmp):
    before = project_files()
    assert client.post("/api/insights", json={"rows": statement()}).status_code == 200
    assert post_csv(client, SAMPLE.read_bytes()).status_code == 200
    assert project_files() - before == set()
    assert list(isolated_tmp.iterdir()) == []


def test_insights_logs_hold_no_merchants_or_amounts(client, caplog):
    caplog.set_level(logging.DEBUG)
    assert client.post("/api/insights", json={"rows": statement()}).status_code == 200
    assert post_csv(client, SAMPLE.read_bytes()).status_code == 200

    text = "\n".join(r.getMessage() for r in caplog.records)
    assert "POST /api/insights 200" in text and "rows=8" in text  # the logger did run
    for secret in ["Warung", "WARUNG", "Toko Serba", "BUDI", "Budi", "KOPI KENANGAN", "Kopi Kenangan",
                   "1500000", "1.500.000", "55000", "10000000", "7910000"]:
        assert secret not in text


# ---- reports ---------------------------------------------------------------


def test_pdf_has_highlights(client):
    resp = client.post("/api/report", json={"rows": statement()}, params={"format": "pdf"})
    assert resp.status_code == 200
    with pdfplumber.open(io.BytesIO(resp.content)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "Highlights" in text
    assert "50/30/20 view" in text
    assert "Top merchants by total" in text and "Toko Serba Ada" in text
    assert DISCLAIMER in text


def test_xlsx_insights_sheet(client):
    resp = client.post("/api/report", json={"rows": statement()}, params={"format": "xlsx"})
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.content))["Insights"]
    cells = {r[0]: r[1:] for r in ws.iter_rows(values_only=True) if r and r[0] is not None}
    assert DISCLAIMER in cells
    assert cells["Money in (IDR)"][0] == 10_000_000
    assert cells["Living spend (IDR)"][0] == 2_090_000
    assert cells["Transfers out (IDR)"][0] == 2_000_000
    # Target, Actual, Difference (per month), Share, Guideline
    assert cells["Needs (bills + transport)"][:5] == (5_000_000, 535_000, 4_465_000, 0.0535, 0.5)
    assert cells["Wants (food + shopping)"][1] == 1_555_000
    assert cells["Toko Serba Ada"][:2] == (1, 1_500_000)
    assert cells["Small spends total (IDR)"][0] == 90_000  # every living row under Rp 50.000, the fee too
    assert cells["Fee count"][0] == 1 and cells["Fees total (IDR)"][0] == 15_000
    for section in ("Cash flow", "50/30/20 view", "Top merchants by total", "Small spends", "Fees", "Highlights"):
        assert section in cells
    money_in_row = [c.value for c in ws["A"]].index("Money in (IDR)") + 1
    assert ws.cell(row=money_in_row, column=2).number_format == "#,##0"
    assert not any(c.data_type == "f" for row in ws.iter_rows() for c in row)


def test_xlsx_insights_sheet_neutralizes_injection():
    # Merchant keys are letters only, so injection can't come from real rows; check the sheet writer itself.
    ins = compute_insights(statement())
    ins["top_merchants"]["by_total"][0]["merchant"] = INJECTION
    ins["highlights"][0]["text"] = "+62 " + INJECTION
    rows = [Row(**r) for r in statement()]
    from app.main import summarize
    content = report_xlsx(rows, summarize(rows), ins)
    ws = load_workbook(io.BytesIO(content))["Insights"]
    texts = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert "'" + INJECTION in texts
    assert "'+62 " + INJECTION in texts
    assert not any(c.data_type == "f" for row in ws.iter_rows() for c in row)
